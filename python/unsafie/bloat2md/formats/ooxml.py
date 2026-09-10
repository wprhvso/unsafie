import io
from typing import Final

import mammoth
from pptx import Presentation
from pptx.exc import PackageNotFoundError

from unsafie.bloat2md.domain import ConversionError, Payload
from unsafie.bloat2md.formats.html import to_markdown
from unsafie.bloat2md.markdown import table, trim
from unsafie.bloat2md.sanitize import clean

_STYLES: Final = (
    "p[style-name='Title'] => h1:fresh",
    "p[style-name='Subtitle'] => h2:fresh",
    "p[style-name='Quote'] => blockquote:fresh",
)
_STYLE_MAP: Final = "\n".join(_STYLES)
_PLACEHOLDER: Final = mammoth.images.img_element(lambda image: {"alt": "embedded image"})


def docx(raw: bytes) -> Payload:
    try:
        result = mammoth.convert_to_html(
            io.BytesIO(raw), style_map=_STYLE_MAP, convert_image=_PLACEHOLDER,
        )
    except (OSError, ValueError, KeyError) as error:
        msg = "the document could not be read"
        raise ConversionError(msg) from error

    markdown, dropped = to_markdown(result.value)
    return Payload(markdown=markdown, dropped={"hidden_nodes": dropped} if dropped else {})


def _shape_text(shape: object) -> str:
    frame = getattr(shape, "text_frame", None)
    if frame is None:
        return ""
    return "\n".join(paragraph.text for paragraph in frame.paragraphs).strip()


def _shape_table(shape: object) -> str:
    tbl = getattr(shape, "table", None)
    if tbl is None:
        return ""
    rows = [[cell.text for cell in row.cells] for row in tbl.rows]
    return table(trim(rows))


def pptx(raw: bytes) -> Payload:
    try:
        deck = Presentation(io.BytesIO(raw))
    except (PackageNotFoundError, OSError, ValueError, KeyError) as error:
        msg = "the presentation could not be read"
        raise ConversionError(msg) from error

    sections: list[str] = []
    for number, slide in enumerate(deck.slides, start=1):
        blocks = [
            block for shape in slide.shapes if (block := _shape_table(shape) or _shape_text(shape))
        ]
        notes = ""
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = (frame.text or "").strip() if frame is not None else ""
        if notes:
            blocks.append(f"**Notes:** {notes}")
        if blocks:
            sections.append(f"## Slide {number}\n\n" + "\n\n".join(blocks))

    return Payload(markdown=clean("\n\n".join(sections)).text, pages=len(deck.slides))
